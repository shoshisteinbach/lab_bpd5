# -*- coding: utf-8 -*-
"""
בניית המצגת של שבוע 1 (6 שקפים, עברית).
הרצה (מתוך תיקיית lab):  week1/.venv/bin/python week1/build_deck.py
התוצר: week1/week1_presentation.pptx

הגרפים נלקחים מ-week1/output/visuals (שלב 4), והטבלה מ-week1/output/03_empirical_clean.csv.
"""
import csv
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

VIS = "week1/output/visuals"
FONT = "Arial"

# --- צבעים (אותם צבעים כמו בגרפים) ---
NAVY = "1F3864"     # צבע ראשי
INK = "1F1F1F"      # טקסט
INK2 = "52514E"     # טקסט משני
TINT = "EEF2F8"     # רקע לכרטיסים
LINE = "D9DEE8"     # קווי טבלה
ORANGE = "EB6834"   # אחוזונים נמוכים / הדגשה
BLUE = "2A78D6"     # אחוזונים גבוהים
WHITE = "FFFFFF"


def rgb(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------------------
# פונקציות עזר
# ---------------------------------------------------------------------------
def set_rtl(paragraph, align=PP_ALIGN.RIGHT):
    paragraph.alignment = align
    paragraph._p.get_or_add_pPr().set("rtl", "1")


def style_run(run, size, color=INK, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    rpr.set("lang", "he-IL")
    cs = rpr.find(qn("a:cs"))
    if cs is None:
        cs = rpr.makeelement(qn("a:cs"), {})
        rpr.append(cs)
    cs.set("typeface", FONT)


def text(slide, x, y, w, h, paragraphs, size=16, color=INK, bold=False,
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.TOP, rtl=True, space_after=6, name=None):
    """תיבת טקסט. paragraphs = רשימה של מחרוזות; **כוכביות** = הדגשה."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(paragraphs, str):
        paragraphs = [paragraphs]
    for i, ptext in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if rtl:
            set_rtl(p, align)
        else:
            p.alignment = align
        p.space_after = Pt(space_after)
        for j, part in enumerate(ptext.split("**")):
            if part:
                style_run(p.add_run(), size, color, bold or j % 2 == 1)
                p.runs[-1].text = part
    return box


def rect(slide, x, y, w, h, fill, rounded=True, name=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        s.adjustments[0] = 0.06
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    s.line.fill.background()
    s.shadow.inherit = False
    if name:
        s.name = name
    return s


def circle_number(slide, x, y, d, label, fill=NAVY, color=WHITE, size=14):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = rgb(fill)
    c.line.fill.background()
    c.shadow.inherit = False
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    style_run(p.add_run(), size, color, True)
    p.runs[0].text = label
    return c


def title(slide, t, subtitle=None):
    """כותרת השקף - תמיד באותו מקום ובאותו גודל."""
    text(slide, 0.6, 0.4, 12.13, 0.7, t, size=32, color=NAVY, bold=True,
         anchor=MSO_ANCHOR.MIDDLE, name="Title")
    if subtitle:
        text(slide, 0.6, 1.1, 12.13, 0.45, subtitle, size=16, color=INK2, name="Subtitle")


def picture(slide, path, x, y, w):
    return slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))


def background(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


# ---------------------------------------------------------------------------
# המצגת
# ---------------------------------------------------------------------------
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]

# ===== שקף 1: פתיחה =====
s = prs.slides.add_slide(BLANK)
background(s, NAVY)
text(s, 0.9, 1.0, 11.5, 0.5, "מעבדה בסטטיסטיקה  |  שבוע 1", size=18, color="CADCFC")
text(s, 0.9, 1.75, 11.5, 1.9, ["עקומות גדילה לקוטר ראש העובר (BPD)"],
     size=44, color=WHITE, bold=True, anchor=MSO_ANCHOR.TOP)
text(s, 0.9, 2.75, 11.5, 0.8, "חישוב אחוזונים מהיר בעזרת AI", size=26, color=WHITE)
# שלוש סטטיסטיקות מרכזיות
stats = [("61,997", "בדיקות אולטרסאונד"), ("14 עד 42", "שבועות הריון"), ("5", "אחוזונים: 3, 5, 10, 95, 97.5")]
for i, (num, lab) in enumerate(stats):
    x = 12.43 - 3.6 * (i + 1) - 0.25 * i + 0.25
    text(s, x, 4.15, 3.6, 0.8, num, size=40, color=WHITE, bold=True, rtl=False)
    text(s, x, 4.95, 3.6, 0.5, lab, size=15, color="CADCFC")
# מקום ריק לשמות
rect(s, 0.9, 6.0, 11.53, 0.75, "2B4A7E", name="Names box")
text(s, 1.2, 6.0, 11.0, 0.75, "מגישים/ות:  _______________________________________________",
     size=18, color=WHITE, anchor=MSO_ANCHOR.MIDDLE, name="Names")
notes(s, "שאלת המחקר: לפי העקומות הקיימות יותר מדי עוברים מסומנים כחריגים. "
         "בשבוע הראשון נתנו ל-AI לחשב את האחוזונים של BPD לפי גיל ההריון, "
         "כדי להשוות בהמשך הקורס לניתוח מבוקר יותר.")

# ===== שקף 2: הכלי והפרומפט =====
s = prs.slides.add_slide(BLANK)
background(s, WHITE)
title(s, "הכלי והפרומפט", "עבדנו עם ה-AI בשיחה, שלב אחרי שלב, ואישרנו כל שלב לפני המעבר לבא")
# כרטיס הכלי (ימין)
rect(s, 8.33, 1.85, 4.4, 4.75, NAVY, name="Tool card")
text(s, 8.68, 2.1, 3.7, 0.4, "כלי ה-AI", size=14, color="CADCFC")
text(s, 8.68, 2.5, 3.7, 0.6, "Claude Code", size=28, color=WHITE, bold=True)
text(s, 8.68, 3.15, 3.7, 1.4, ["מודל Claude Opus 5.5", "הקוד נכתב והורץ ב-R",
                               "נתונים: גיל הריון ו-BPD בלבד"], size=15, color=WHITE)
rect(s, 8.58, 4.85, 3.9, 1.5, "2B4A7E", name="Prompt quote")
text(s, 8.8, 4.85, 3.46, 1.5,
     "\"קרא את כל הקבצים, בנה תוכנית מורכבת מכמה שלבים, שאל אותי במקום לנחש, "
     "ואל תתחיל עד שאאשר\"", size=14, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
# מהלך השיחה (שמאל) - 6 צעדים
steps = [
    ("הפרומפט הראשון", "הדבקנו את הנחיות המטלה וביקשנו תוכנית עבודה ושאלות"),
    ("בחירת מדד", "BPD, ובדיקה שהתוכנית מתאימה ל-8 שלבי הקורס"),
    ("אופן העבודה", "קוד R לכל שלב, הצגת ממצאים והמתנה לאישור"),
    ("ניקוי", "הסרת שגיאות רישום, ושמירה גם של הנתונים הגולמיים"),
    ("חישוב", "אחוזונים אמפיריים לכל שבוע, 14 עד 42"),
    ("הצגה", "גרפים מעוצבים, מסמך סיכום ומצגת"),
]
for i, (head, body) in enumerate(steps):
    y = 1.9 + i * 0.8
    circle_number(s, 7.4, y + 0.05, 0.5, str(i + 1))
    text(s, 0.6, y, 6.6, 0.35, head, size=16, color=INK, bold=True)
    text(s, 0.6, y + 0.35, 6.6, 0.4, body, size=14, color=INK2)
notes(s, "בנוסף: ה-AI הוסיף על דעת עצמו השוואה לעקומת Hadlock. ביקשנו לא לעשות מה שלא התבקש, "
         "וההשוואה הוסרה.")

# ===== שקף 3: הנתונים והניקוי =====
s = prs.slides.add_slide(BLANK)
background(s, WHITE)
title(s, "הנתונים: שגיאות רישום שחייבים לנקות",
      "בסקאלה לוגריתמית רואים שלושה \"פסים\" במקום אחד")
callouts = [("1.4%", "נרשמו בס\"מ במקום במ\"מ", "889 ערכים, למשל 8.5 במקום 85", ORANGE),
            ("0.4%", "ספרה מיותרת", "244 ערכים, למשל 892 במקום 89.2", BLUE),
            ("98.1%", "נשארו אחרי הניקוי", "60,799 בדיקות", NAVY)]
for i, (num, head, body, col) in enumerate(callouts):
    y = 1.85 + i * 1.65
    rect(s, 9.13, y, 3.6, 1.45, TINT, name=f"Callout {i+1}")
    text(s, 9.38, y + 0.12, 3.1, 0.65, num, size=32, color=col, bold=True)
    text(s, 9.38, y + 0.75, 3.1, 0.3, head, size=15, color=INK, bold=True)
    text(s, 9.38, y + 1.06, 3.1, 0.3, body, size=12, color=INK2)
picture(s, f"{VIS}/04_raw_errors.png", 0.5, 1.75, 8.4)
text(s, 0.6, 6.45, 8.2, 0.6, "בלי ניקוי, אחוזון 3 יוצא נמוך ב-1 עד 3 מ\"מ, "
     "כי רוב השגיאות הן ערכים נמוכים מדי", size=14, color=INK2)
notes(s, "כלל 1: הסרת ערכים מחוץ לטווח 10-120 מ\"מ (1,133 ערכים). "
         "כלל 2: הסרת ערכים שרחוקים מהחציון של השבוע ביותר מ-5 סטיות תקן חסינות (65 ערכים).")

# ===== שקף 4: שיטת החישוב =====
s = prs.slides.add_slide(BLANK)
background(s, WHITE)
title(s, "איך ה-AI חישב את העקומות",
      "אחוזונים אמפיריים: ישירות מהנתונים הנקיים, לכל שבוע בנפרד")
rect(s, 6.83, 1.85, 5.9, 3.95, TINT, name="Method card")
method_steps = [
    "לוקחים את כל הבדיקות של שבוע מסוים, למשל כל הבדיקות בשבוע 20",
    "מסדרים אותן מהקטנה לגדולה",
    "אחוזון 3 = הערך ש-3% מהבדיקות נמצאות מתחתיו. כך גם 5, 10, 95 ו-97.5",
]
for i, step in enumerate(method_steps):
    y = 2.15 + i * 0.85
    circle_number(s, 11.93, y, 0.5, str(i + 1))
    text(s, 7.13, y - 0.02, 4.6, 0.75, step, size=15, color=INK, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 7.13, 4.8, 5.3, 0.65, WHITE, name="Formula")
text(s, 7.13, 4.8, 5.3, 0.65, "P3 = quantile( BPD[week] , 0.03 )", size=16, color=INK,
     rtl=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
picture(s, f"{VIS}/04_n_per_week.png", 0.45, 2.0, 6.1)
text(s, 0.6, 4.6, 5.9, 0.6, "מספר הבדיקות בכל שבוע. בשבועות 43 ו-44 כמעט אין נתונים",
     size=13, color=INK2)
text(s, 0.6, 6.05, 12.13, 0.9,
     ["**למה עד שבוע 42?** בשבועות 43 ו-44 יש רק 12 בדיקות, ואי אפשר לחשב מהן אחוזון 3. "
      "בשבוע 42 יש 153 בדיקות, ולכן האחוזונים הקיצוניים שם פחות יציבים."],
     size=15, color=INK2)
notes(s, "אין כאן שום הנחה על צורת ההתפלגות: כל אחוזון נלקח ישירות מהבדיקות של אותו שבוע. "
         "החישוב נעשה בפונקציה quantile של R.")

# ===== שקף 5: התוצאות =====
s = prs.slides.add_slide(BLANK)
background(s, WHITE)
title(s, "התוצאה: עקומות האחוזונים של BPD",
      "כתום: אחוזונים נמוכים (חשד לעובר קטן). כחול: אחוזונים גבוהים (חשד לעובר גדול)")
picture(s, f"{VIS}/04_growth_chart.png", 0.45, 1.7, 8.0)
# טבלה: שבועות נבחרים מהאחוזונים האמפיריים
with open("week1/output/03_empirical_clean.csv", encoding="utf-8") as f:
    rows = [r for r in csv.DictReader(f) if int(r["week"]) in (16, 20, 24, 28, 32, 36, 40)]
cols = ["week", "P3", "P5", "P10", "P95", "P97.5"]
heads = ["שבוע", "3", "5", "10", "95", "97.5"]
tbl_shape = s.shapes.add_table(len(rows) + 1, len(cols), Inches(8.73), Inches(1.95),
                               Inches(4.0), Inches(3.9))
tbl_shape.name = "Percentile table"
tbl = tbl_shape.table
tbl._tbl.tblPr.set("rtl", "1")
tbl._tbl.tblPr.set("bandRow", "0")
tbl._tbl.tblPr.set("firstRow", "0")
widths = [0.8, 0.64, 0.64, 0.64, 0.64, 0.64]
for j, w in enumerate(widths):
    tbl.columns[j].width = Inches(w)
head_colours = [NAVY, "C4471A", ORANGE, "D9825B", "4F8FDB", "1C5CAB"]
for i in range(len(rows) + 1):
    for j, c in enumerate(cols):
        cell = tbl.cell(i, j)
        cell.margin_left = cell.margin_right = Inches(0.04)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        if i == 0:
            cell.fill.fore_color.rgb = rgb(head_colours[j])
            val, colr, bold = heads[j], WHITE, True
        else:
            cell.fill.fore_color.rgb = rgb(WHITE if i % 2 else TINT)
            v = float(rows[i - 1][c])
            val = str(int(v)) if v == int(v) else f"{v:.1f}"
            colr, bold = (NAVY, True) if c == "week" else (INK, False)
        p = cell.text_frame.paragraphs[0]
        set_rtl(p, PP_ALIGN.CENTER)
        style_run(p.add_run(), 14, colr, bold)
        p.runs[0].text = val
text(s, 8.73, 1.55, 4.0, 0.35, "אחוזונים (מ\"מ), שבועות נבחרים", size=13, color=INK2, bold=True)
text(s, 8.73, 6.0, 4.0, 0.9, "הטבלה המלאה לכל שבוע נמצאת במסמך הסיכום",
     size=12, color=INK2)
notes(s, "בתחום האפור (בין אחוזון 10 לאחוזון 95) העובר אינו מסומן כחריג. "
         "האחוזונים כמעט תמיד מספרים שלמים, כי המדידות נרשמו במ\"מ שלמים.")

# ===== שקף 6: על מה צריך להיזהר =====
s = prs.slides.add_slide(BLANK)
background(s, WHITE)
title(s, "מה למדנו: על מה צריך להיזהר בעבודה עם AI")
lessons = [
    ("לבדוק את הנתונים לפני החישוב",
     "1.8% מהערכים נרשמו ביחידות שגויות. חישוב ישיר היה מוריד את אחוזון 3 בעד 3 מ\"מ"),
    ("לוודא מה ה-AI עשה בפועל",
     "הוא הוסיף השוואה שלא התבקשה, וזיהה בטעות נוסחה אחרת כנוסחת Hadlock. ההשוואה הוסרה"),
    ("שבועות עם מעט נתונים",
     "בשבוע 42 האחוזונים הקיצוניים נשענים על 4 עד 5 בדיקות בלבד, וה-AI לא סימן זאת מעצמו"),
    ("מה עוד לא נבדק",
     "מדידות חוזרות של אותה אישה, חלוקה לאימון ומבחן, והשפעת הטכנאי. בהמשך הקורס"),
]
for i, (head, body) in enumerate(lessons):
    y = 1.5 + i * 1.33
    rect(s, 6.83, y, 5.9, 1.15, TINT, name=f"Lesson {i+1}")
    circle_number(s, 12.03, y + 0.3, 0.5, str(i + 1), fill=ORANGE)
    text(s, 7.08, y + 0.14, 4.75, 0.35, head, size=16, color=INK, bold=True)
    text(s, 7.08, y + 0.5, 4.75, 0.6, body, size=13, color=INK2)
picture(s, f"{VIS}/04_cleaning_effect_P3.png", 0.45, 2.3, 6.2)
text(s, 0.6, 5.0, 6.0, 0.8,
     ["**השפעת הניקוי:** בכמה מ\"מ אחוזון 3 יורד בכל שבוע אם מחשבים על הנתונים "
      "הגולמיים, בלי ניקוי."], size=13, color=INK2)
notes(s, "המטרה בשבוע הזה הייתה לראות את התוצאה של עבודה מהירה עם AI. "
         "בהמשך הקורס נשווה לניתוח מבוקר יותר.")

prs.save("week1/week1_presentation.pptx")
print("Saved week1/week1_presentation.pptx")
